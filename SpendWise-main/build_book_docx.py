# -*- coding: utf-8 -*-
"""
בונה את ספר הפרויקט של "שקלול" (SpendWise) כקובץ Word.
הרצה:  .venv/Scripts/python.exe build_book_docx.py
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "ספר פרויקט שקלול.docx")

HEBREW_FONT = "David"
CODE_FONT = "Consolas"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x55, 0x55, 0x55)


# ─────────────────────────── helpers ───────────────────────────

def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return paragraph


def style_font(style, name=HEBREW_FONT, size=None, bold=None, color=None):
    style.font.name = name
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def setup(doc):
    for s in doc.sections:
        s.top_margin = Cm(2.2)
        s.bottom_margin = Cm(2.2)
        s.left_margin = Cm(2.4)
        s.right_margin = Cm(2.4)
        # RTL section
        sectPr = s._sectPr
        rtlGutter = OxmlElement("w:rtlGutter")
        sectPr.append(rtlGutter)
        bidi = OxmlElement("w:bidi")
        sectPr.append(bidi)

    normal = doc.styles["Normal"]
    style_font(normal, HEBREW_FONT, 11.5)

    for name, size, color in (("Heading 1", 20, ACCENT),
                              ("Heading 2", 16, ACCENT),
                              ("Heading 3", 13.5, ACCENT)):
        st = doc.styles[name]
        style_font(st, HEBREW_FONT, size, True, color)
        st.paragraph_format.space_before = Pt(16 if name == "Heading 1" else 12)
        st.paragraph_format.space_after = Pt(8)

    for name in ("List Bullet", "List Number"):
        try:
            style_font(doc.styles[name], HEBREW_FONT, 11.5)
        except KeyError:
            pass


def H(doc, text, level=1, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    r.font.name = HEBREW_FONT
    rpr = r._element.get_or_add_rPr()
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), HEBREW_FONT)
    rpr.append(rfonts)
    r.font.size = Pt({1: 20, 2: 16, 3: 13.5}[level])
    r.font.bold = True
    r.font.color.rgb = ACCENT
    set_rtl(p)
    return p


def P(doc, text, bold=False, size=11.5, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = HEBREW_FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    rpr = r._element.get_or_add_rPr()
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), HEBREW_FONT)
    rpr.append(rfonts)
    set_rtl(p)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def rich(doc, parts, space_after=6):
    """parts = list of (text, bold) tuples."""
    p = doc.add_paragraph()
    for text, bold in parts:
        r = p.add_run(text)
        r.font.name = HEBREW_FONT
        r.font.size = Pt(11.5)
        r.font.bold = bold
        rpr = r._element.get_or_add_rPr()
        rfonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), HEBREW_FONT)
        rpr.append(rfonts)
    set_rtl(p)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def BULLET(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.font.name = HEBREW_FONT
        r.font.size = Pt(11.5)
        r.font.bold = True
        rpr = r._element.get_or_add_rPr()
        rfonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), HEBREW_FONT)
        rpr.append(rfonts)
    r = p.add_run(text)
    r.font.name = HEBREW_FONT
    r.font.size = Pt(11.5)
    rpr = r._element.get_or_add_rPr()
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), HEBREW_FONT)
    rpr.append(rfonts)
    set_rtl(p)
    return p


def CODE(doc, code, caption=None):
    if caption:
        P(doc, caption, bold=True, size=10.5, color=MUTED, space_after=2)
    for line in code.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        r = p.add_run(line if line.strip() else " ")
        r.font.name = CODE_FONT
        r.font.size = Pt(9.5)
        rpr = r._element.get_or_add_rPr()
        rfonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), CODE_FONT)
        rpr.append(rfonts)
        # LTR for code
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "0")
        pPr.append(bidi)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = Pt(11)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def TABLE(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        r.font.name = HEBREW_FONT
        r.font.bold = True
        r.font.size = Pt(10.5)
        rpr = r._element.get_or_add_rPr()
        rfonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), HEBREW_FONT)
        rpr.append(rfonts)
        set_rtl(p)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(val))
            r.font.name = HEBREW_FONT
            r.font.size = Pt(10.5)
            rpr = r._element.get_or_add_rPr()
            rfonts = OxmlElement("w:rFonts")
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                rfonts.set(qn(attr), HEBREW_FONT)
            rpr.append(rfonts)
            set_rtl(p)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t
